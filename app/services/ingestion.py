from pathlib import Path
from typing import Iterable
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from docx import Document as DocxDocument

VALID_EXTENSIONS = [".txt", ".pdf", ".docx", ".md"]

def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    if suffix not in VALID_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")
    
    try: 
        if suffix == ".pdf":
            return PyPDFLoader(str(path)).load()
        elif suffix == ".docx":
            doc = DocxDocument(path)
            text = "\n".join([para.text for para in doc.paragraphs])
            return [Document(page_content=text, metadata={"source": str(path)})]
        elif suffix in [".txt", ".md"]:
            with open(path, "r", encoding="utf-8") as f:
                text = f.read()
            return [Document(page_content=text, metadata={"source": str(path)})]
    except Exception as e:
        raise ValueError(f"Error loading file {path}: {e}")

def chunk_documents(documents: Iterable[Document]) -> list[Document]:
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=125, add_start_index=True)
    return text_splitter.split_documents(documents)

